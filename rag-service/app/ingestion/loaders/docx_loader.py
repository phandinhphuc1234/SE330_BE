from pathlib import Path

from docx import Document

from app.ingestion.loaders.base import BaseLoader, RawDocument


class DocxLoader(BaseLoader):
    supported_extensions = {".docx"}

    def load(self, file_path: Path) -> list[RawDocument]:
        document = Document(file_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return [RawDocument(text=text, metadata={"source": str(file_path)})]
