from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document as DocxDocument

from app.ingestion.parsers.base import DocumentParser, ParsedDocument


class MarkdownParser(DocumentParser):
    """Parse Markdown files as a single document unit."""

    name = "markdown"
    supported_extensions = {".md", ".markdown"}

    def parse(self, file_path: Path) -> list[ParsedDocument]:
        return [
            ParsedDocument(
                text=file_path.read_text(encoding="utf-8"),
                metadata={
                    "source": str(file_path),
                    "parser": self.name,
                },
            )
        ]


class HtmlParser(DocumentParser):
    """Parse HTML into plain text using BeautifulSoup."""

    name = "html"
    supported_extensions = {".html", ".htm"}

    def parse(self, file_path: Path) -> list[ParsedDocument]:
        soup = BeautifulSoup(file_path.read_text(encoding="utf-8"), "lxml")
        return [
            ParsedDocument(
                text=soup.get_text("\n"),
                metadata={
                    "source": str(file_path),
                    "parser": self.name,
                },
            )
        ]


class DocxParser(DocumentParser):
    """Parse DOCX paragraphs as a single document unit."""

    name = "docx"
    supported_extensions = {".docx"}

    def parse(self, file_path: Path) -> list[ParsedDocument]:
        document = DocxDocument(file_path)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        return [
            ParsedDocument(
                text=text,
                metadata={
                    "source": str(file_path),
                    "parser": self.name,
                },
            )
        ]
